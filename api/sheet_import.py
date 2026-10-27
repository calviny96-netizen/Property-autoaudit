"""Read an uploaded CSV or Excel file into plain rows of text, for import previews.

Nothing is stored. The browser shows the rows, lets the person pick the sheet and the
columns, and saves the result through the normal endpoints (the sales watchlist uses
PUT /stock/tracked).

Every cell comes back as a string. Numbers are written out in full, the way people
typed them: Excel keeps a phone column as numbers (6281234567890.0) or, in a CSV saved
from a narrow column, as scientific notation (8.12345E+11); both become plain digits.
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000          # rows returned per sheet
MAX_SCAN_ROWS = 100_000  # rows counted per sheet (a guard against huge workbooks)
MAX_SHEETS = 30
MAX_COLUMNS = 100
MAX_CELL = 300
CSV_DELIMITERS = (',', ';', '\t')
EXCEL_SUFFIXES = ('.xlsx', '.xlsm')

TOO_LARGE = 'The file is too large. Use a file of at most 5 MB.'
WRONG_TYPE = 'Please choose a CSV or Excel file (.csv or .xlsx).'
OLD_EXCEL = 'Old Excel files (.xls) cannot be read. In Excel, choose File > Save As > Excel Workbook (.xlsx) or CSV, then upload that file.'
UNREADABLE_EXCEL = 'This Excel file could not be read. Save it again as .xlsx (without a password) or as CSV, then try again.'
NO_ROWS = 'This file has no rows to import.'

_SCIENTIFIC = re.compile(r'[+-]?\d+(?:[.,]\d+)?[eE][+-]?\d+')
_WHOLE_DECIMAL = re.compile(r'[+-]?\d+\.0+')
_EXCEL_TEXT_FORMULA = re.compile(r'="(.*)"')  # ="081234567890" keeps the leading zero in a CSV


class SheetImportError(ValueError):
    """A problem with the uploaded file, worded for the person who chose it."""


def number_text(value):
    """``6281234567890.0`` -> ``'6281234567890'``, ``'8.12345E+11'`` -> ``'812345000000'``, ``1.5`` -> ``'1.5'``."""
    try:
        amount = value if isinstance(value, Decimal) else Decimal(str(value).replace(',', '.'))
    except InvalidOperation:
        return str(value)
    if not amount.is_finite():
        return str(value)
    if amount == amount.to_integral_value():
        return str(int(amount))
    return format(amount.normalize(), 'f')


def cell_text(value):
    """One cell as trimmed text (numbers in full, dates as ISO, at most MAX_CELL characters)."""
    if value is None:
        return ''
    if isinstance(value, bool):
        text = 'TRUE' if value else 'FALSE'
    elif isinstance(value, (int, float, Decimal)):
        text = number_text(value)
    elif isinstance(value, datetime):
        text = value.date().isoformat() if value.time() == time.min else value.isoformat(sep=' ', timespec='minutes')
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    else:
        text = str(value).replace('\r\n', '\n').replace('\r', '\n').strip()
        quoted = _EXCEL_TEXT_FORMULA.fullmatch(text)
        if quoted:
            text = quoted.group(1).strip()
        if _SCIENTIFIC.fullmatch(text) or _WHOLE_DECIMAL.fullmatch(text):
            text = number_text(text)
    return text.strip()[:MAX_CELL]


def clean_row(values):
    """Cells as text without the empty cells at the end; an empty list for a blank row."""
    row = [cell_text(value) for value in list(values)[:MAX_COLUMNS]]
    while row and not row[-1]:
        row.pop()
    return row


def collect(name, rows):
    """The first MAX_ROWS non-empty rows of a sheet and how many non-empty rows it has."""
    kept, total = [], 0
    for values in rows:
        row = clean_row(values)
        if not row:
            continue
        total += 1
        if len(kept) < MAX_ROWS:
            kept.append(row)
        if total >= MAX_SCAN_ROWS:
            break
    return {'name': str(name)[:100], 'rows': kept, 'total_rows': total}


# --------------------------------------------------------------------------- CSV
def decode_text(data):
    """UTF-8 (with or without a byte-order mark), UTF-16 with a mark, else Windows-1252 / Latin-1."""
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        try:
            return data.decode('utf-16')
        except UnicodeDecodeError:
            pass
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode('latin-1')


def detect_delimiter(text):
    """Comma, semicolon or tab: the one that splits the most of the first lines into the same number of columns."""
    sample = '\n'.join([line for line in text.splitlines() if line.strip()][:50])
    best, best_score = ',', (0, 0)
    for delimiter in CSV_DELIMITERS:
        try:
            widths = [len(row) for row in csv.reader(io.StringIO(sample), delimiter=delimiter) if len(row) > 1]
        except csv.Error:
            continue
        if not widths:
            continue
        common = max(set(widths), key=lambda width: (widths.count(width), width))
        score = (widths.count(common), common)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def read_csv(data, name):
    text = decode_text(data).replace('\x00', '')
    try:
        rows = csv.reader(io.StringIO(text), delimiter=detect_delimiter(text))
        return [collect(name, rows)]
    except csv.Error as exc:
        raise SheetImportError(f'This CSV file could not be read ({exc}). Save it again as CSV and try again.') from exc


# --------------------------------------------------------------------------- Excel
def read_excel(data):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError, TypeError) as exc:
        raise SheetImportError(UNREADABLE_EXCEL) from exc
    try:
        visible, hidden = [], []
        for sheet in book.worksheets[:MAX_SHEETS]:
            sheet.reset_dimensions()  # some programs store a wrong size; read every row that is there
            result = collect(sheet.title, sheet.iter_rows(values_only=True))
            if result['total_rows']:
                (visible if getattr(sheet, 'sheet_state', 'visible') == 'visible' else hidden).append(result)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError) as exc:
        raise SheetImportError(UNREADABLE_EXCEL) from exc
    finally:
        book.close()
    # Hidden helper sheets are left out unless they are the only ones with data.
    return visible or hidden


# --------------------------------------------------------------------------- entry point
def base_name(file_name):
    return re.split(r'[\\/]', file_name or '')[-1].strip()[:200]


def read_sheets(file_name, data):
    """``{'file_name', 'sheets': [{'name', 'rows', 'total_rows'}]}`` for a .csv, .xlsx or .xlsm upload.

    Empty rows are skipped, empty cells at the end of a row are dropped, and only sheets
    with at least one row are returned. Raises SheetImportError with a readable message.
    """
    name = base_name(file_name)
    lower = name.lower()
    if lower.endswith('.xls'):
        raise SheetImportError(OLD_EXCEL)
    if not lower.endswith(('.csv', *EXCEL_SUFFIXES)):
        raise SheetImportError(WRONG_TYPE)
    if len(data) > MAX_BYTES:
        raise SheetImportError(TOO_LARGE)
    if not data.strip():
        raise SheetImportError('The file is empty.')
    if data.startswith(b'\xd0\xcf\x11\xe0'):  # an old .xls or a password-protected workbook
        raise SheetImportError(OLD_EXCEL if lower.endswith('.csv') else UNREADABLE_EXCEL)
    # A workbook saved with a .csv name is still read as a workbook.
    if lower.endswith(EXCEL_SUFFIXES) or data.startswith(b'PK\x03\x04'):
        sheets = read_excel(data)
    else:
        sheets = read_csv(data, name.rsplit('.', 1)[0] or 'Sheet 1')
    sheets = [sheet for sheet in sheets if sheet['total_rows']]
    if not sheets:
        raise SheetImportError(NO_ROWS)
    return {'file_name': name, 'sheets': sheets}
