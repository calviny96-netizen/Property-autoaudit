"""CSV / Excel reading for "Import from CSV / Excel" (sales watchlist) and its endpoint's access rule."""
import io
import unittest

from harness import ServerTestCase
from sheet_import import MAX_BYTES, MAX_ROWS, SheetImportError, cell_text, detect_delimiter, number_text, read_sheets


def workbook_bytes(build):
    from openpyxl import Workbook
    book = Workbook()
    build(book)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class NumberText(unittest.TestCase):
    def test_numbers_are_written_out_in_full(self):
        self.assertEqual(number_text(6281234567890.0), '6281234567890')
        self.assertEqual(number_text(81234567890), '81234567890')
        self.assertEqual(number_text('8.12345E+11'), '812345000000')
        self.assertEqual(number_text('6,28123E+12'), '6281230000000')
        self.assertEqual(number_text(1.5), '1.5')
        self.assertEqual(number_text(1e-7), '0.0000001')

    def test_text_cells(self):
        self.assertEqual(cell_text('  Budi Santoso \r\n'), 'Budi Santoso')
        self.assertEqual(cell_text('6281234567890.0'), '6281234567890')
        self.assertEqual(cell_text('8.12345E+11'), '812345000000')
        self.assertEqual(cell_text('="081234567890"'), '081234567890')
        self.assertEqual(cell_text('123,000'), '123,000')  # thousands, not a decimal
        self.assertEqual(cell_text('0812-3456-7890'), '0812-3456-7890')
        self.assertEqual(cell_text(None), '')
        self.assertEqual(len(cell_text('x' * 1000)), 300)


class CsvFiles(unittest.TestCase):
    def rows(self, text, encoding='utf-8', name='sales.csv'):
        result = read_sheets(name, text.encode(encoding) if isinstance(text, str) else text)
        self.assertEqual(len(result['sheets']), 1)
        return result['sheets'][0]['rows']

    def test_delimiters(self):
        expected = [['Nama', 'No HP'], ['Budi', '081234567890'], ['Sari', '6282233744657']]
        self.assertEqual(self.rows('Nama,No HP\nBudi,081234567890\nSari,6282233744657\n'), expected)
        self.assertEqual(self.rows('Nama;No HP\r\nBudi;081234567890\r\nSari;6282233744657\r\n'), expected)
        self.assertEqual(self.rows('Nama\tNo HP\nBudi\t081234567890\nSari\t6282233744657\n'), expected)

    def test_semicolon_file_with_commas_inside_cells(self):
        text = 'Nama;No HP;Alamat\nBudi;0812 3456 7890;Jl. Merdeka, No 5\nSari;"0813 1111 2222, 0813 3333 4444";Jl. Mawar\n'
        self.assertEqual(detect_delimiter(text), ';')
        self.assertEqual(self.rows(text)[2], ['Sari', '0813 1111 2222, 0813 3333 4444', 'Jl. Mawar'])

    def test_single_column_of_numbers(self):
        self.assertEqual(self.rows('081234567890\n6282233744657\n'), [['081234567890'], ['6282233744657']])

    def test_byte_order_mark_and_latin1(self):
        self.assertEqual(self.rows('\ufeffName,Phone\nBudi,081234567890\n')[0], ['Name', 'Phone'])
        self.assertEqual(self.rows('Name,Phone\nJos\xe9,081234567890\n', encoding='latin-1')[1], ['Jos\xe9', '081234567890'])
        self.assertEqual(self.rows('Name\tPhone\nAndi\t0812\n', encoding='utf-16')[1], ['Andi', '0812'])

    def test_scientific_numbers_empty_rows_and_trailing_columns(self):
        text = 'Name,Phone,,\n\n,,,\nBudi,8.12345E+11,,\nSari,6282233744657.0,,\n , ,\n'
        result = read_sheets('sales.csv', text.encode())
        sheet = result['sheets'][0]
        self.assertEqual(sheet['rows'], [['Name', 'Phone'], ['Budi', '812345000000'], ['Sari', '6282233744657']])
        self.assertEqual(sheet['total_rows'], 3)
        self.assertEqual(result['file_name'], 'sales.csv')

    def test_only_the_first_rows_are_returned(self):
        text = 'Phone\n' + ''.join(f'0812000{n:05d}\n' for n in range(MAX_ROWS + 200))
        sheet = read_sheets('many.csv', text.encode())['sheets'][0]
        self.assertEqual(len(sheet['rows']), MAX_ROWS)
        self.assertEqual(sheet['total_rows'], MAX_ROWS + 201)

    def test_file_name_without_folders(self):
        self.assertEqual(read_sheets('C:\\fakepath\\team.csv', b'0812\n')['file_name'], 'team.csv')


class ExcelFiles(unittest.TestCase):
    def test_sheets_numeric_phones_and_empty_rows(self):
        def build(book):
            first = book.active
            first.title = 'Surabaya'
            first.append(['Nama', 'No HP', None])
            first.append(['Budi', 6281234567890, None])
            first.append([None, None, None])
            first.append(['Sari', 6282233744657.0])
            first.append(['Andi', 8.12345e11])
            first.append(['Rina', '0813-1111-2222'])
            first['E2'] = ''  # an empty cell at the end of a row
            second = book.create_sheet('Jakarta')
            second.append(['Agent', 'WhatsApp'])
            second.append(['Dewi', 81202310022])
            book.create_sheet('Empty')
            hidden = book.create_sheet('Lists')
            hidden.append(['helper'])
            hidden.sheet_state = 'hidden'
        result = read_sheets('team.xlsx', workbook_bytes(build))
        self.assertEqual([sheet['name'] for sheet in result['sheets']], ['Surabaya', 'Jakarta'])
        surabaya, jakarta = result['sheets']
        self.assertEqual(surabaya['rows'], [
            ['Nama', 'No HP'], ['Budi', '6281234567890'], ['Sari', '6282233744657'], ['Andi', '812345000000'], ['Rina', '0813-1111-2222']])
        self.assertEqual(surabaya['total_rows'], 5)
        self.assertEqual(jakarta['rows'], [['Agent', 'WhatsApp'], ['Dewi', '81202310022']])
        self.assertEqual(jakarta['total_rows'], 2)

    def test_xlsm_and_a_workbook_saved_as_csv(self):
        data = workbook_bytes(lambda book: book.active.append(['Budi', 6281234567890]))
        for name in ('team.xlsm', 'team.csv'):
            self.assertEqual(read_sheets(name, data)['sheets'][0]['rows'], [['Budi', '6281234567890']])

    def test_only_hidden_sheets_with_data_are_still_read(self):
        def build(book):
            book.active.title = 'Blank'
            hidden = book.create_sheet('Data')
            hidden.append(['0812'])
            hidden.sheet_state = 'hidden'
        self.assertEqual([sheet['name'] for sheet in read_sheets('team.xlsx', workbook_bytes(build))['sheets']], ['Data'])


class Rejections(unittest.TestCase):
    def rejects(self, name, data, message):
        with self.assertRaises(SheetImportError) as caught:
            read_sheets(name, data)
        self.assertIn(message, str(caught.exception))

    def test_types_size_and_broken_files(self):
        self.rejects('sales.txt', b'0812', 'CSV or Excel file')
        self.rejects('sales.pdf', b'%PDF', 'CSV or Excel file')
        self.rejects('sales', b'0812', 'CSV or Excel file')
        self.rejects('sales.xls', b'\xd0\xcf\x11\xe0rest', '.xls')
        self.rejects('sales.csv', b'0' * (MAX_BYTES + 1), 'at most 5 MB')
        self.rejects('sales.csv', b'', 'empty')
        self.rejects('sales.csv', b'\n , ,\n\n', 'no rows')
        self.rejects('sales.xlsx', b'this is not a workbook', 'could not be read')
        self.rejects('sales.xlsx', b'\xd0\xcf\x11\xe0protected', 'could not be read')
        self.rejects('blank.xlsx', workbook_bytes(lambda book: None), 'no rows')


class ParseEndpoint(ServerTestCase):
    def post_file(self, client, name, data, status):
        boundary = 'xm-sheet-import-boundary'
        raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\n'
               f'Content-Type: application/octet-stream\r\n\r\n').encode() + data + f'\r\n--{boundary}--\r\n'.encode()
        return self.call(client, '/stock/tracked/parse', 'POST', raw=raw,
                         content_type='multipart/form-data; boundary=' + boundary, status=status)

    def test_company_admin_only(self):
        _, _, boss = self.make_company('Sheet import')
        _, staff = self.make_member(boss, 'Staff')
        csv_bytes = '\ufeffNama;No HP\nBudi;081234567890\n'.encode()
        result = self.post_file(boss, 'sales.csv', csv_bytes, 200)
        self.assertEqual(result, {'file_name': 'sales.csv', 'sheets': [
            {'name': 'sales', 'rows': [['Nama', 'No HP'], ['Budi', '081234567890']], 'total_rows': 2}]})
        failure = self.post_file(staff, 'sales.csv', csv_bytes, 403)
        self.assertIn('company admin', failure['detail'])
        self.assertIn('CSV or Excel', self.post_file(boss, 'sales.txt', b'0812', 400)['detail'])
        # Reading a file never changes the watchlist.
        self.assertEqual(self.call(boss, '/stock/overview')['tracked'], [])


if __name__ == '__main__':
    unittest.main()
